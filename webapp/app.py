import csv
import io
import os
import re
import sys
import tempfile

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from flask import Flask, jsonify, make_response, render_template, request

from churn_predict import predict_churn, predict_batch, THRESHOLD

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024

UPLOAD_DIR = os.path.join(tempfile.gettempdir(), "churn_uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

REQUIRED = ["Geography", "Gender", "Age"]
NUMBER_FIELDS = {"Age": "float", "Balance": "float", "NumOfProducts": "int", "IsActiveMember": "int"}
DEFAULTS = {"Age": 0.0, "Balance": 0.0, "NumOfProducts": 1, "IsActiveMember": 1}

COLUMN_ALIASES = {
    "geography": "Geography", "country": "Geography",
    "gender": "Gender", "sex": "Gender",
    "age": "Age", "ageyears": "Age",
    "balance": "Balance",
    "numofproducts": "NumOfProducts", "num_products": "NumOfProducts", "products": "NumOfProducts",
    "isactivemember": "IsActiveMember", "activemember": "IsActiveMember", "active": "IsActiveMember",
}
GEO_MAP = {"france": "France", "germany": "Germany", "spain": "Spain",
           "fr": "France", "de": "Germany", "es": "Spain"}
GENDER_MAP = {"male": "Male", "female": "Female", "m": "Male", "f": "Female",
              "ذكر": "Male", "أنثى": "Female"}


def clean_col(c):
    return re.sub(r"[\s\-_]+", "", str(c).strip().lower())


def normalize_table(rows):
    renamed = {}
    if rows:
        for col in rows[0]:
            key = clean_col(col)
            if key in COLUMN_ALIASES:
                renamed[col] = COLUMN_ALIASES[key]
    out = []
    for r in rows:
        out.append({renamed.get(k, k): v for k, v in r.items()})
    return out


def coerce_values(rows):
    errs = []
    for r in rows:
        for col, ctype in NUMBER_FIELDS.items():
            if col not in r:
                r[col] = DEFAULTS[col]
                continue
            try:
                v = float(str(r[col]).replace(",", "").strip())
                r[col] = int(v) if ctype == "int" else v
            except (ValueError, TypeError):
                r[col] = DEFAULTS[col]
                if col not in errs:
                    errs.append(col)
        r["Geography"] = GEO_MAP.get(str(r.get("Geography", "")).strip().lower(), "France")
        r["Gender"] = GENDER_MAP.get(str(r.get("Gender", "")).strip().lower(), "Male")
    return rows, errs


def read_table(path, filename):
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xls"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        rows = []
        for row in it:
            if any(v is not None and v != "" for v in row):
                rows.append(dict(zip(header, row)))
        wb.close()
        return rows
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                rows = [dict(r) for r in csv.DictReader(f)]
            rows = [r for r in rows if any(v is not None and v != "" for v in r.values())]
            return rows
        except (UnicodeDecodeError, csv.Error):
            continue
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        return [dict(r) for r in csv.DictReader(f)]


def rows_to_json(rows, columns):
    out = []
    for r in rows:
        out.append({c: r.get(c) for c in columns})
    return out


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/health")
def health():
    return jsonify({"status": "ok", "threshold": THRESHOLD})


@app.route("/api/predict", methods=["POST"])
def api_predict():
    data = request.get_json(force=True)
    try:
        res = predict_churn(
            geography=data["geography"],
            gender=data["gender"],
            age=float(data["age"]),
            balance=float(data.get("balance", 0.0)),
            num_of_products=int(data.get("num_of_products", 1)),
            is_active_member=int(data.get("is_active_member", 1)),
        )
    except KeyError as e:
        return jsonify({"error": f"Missing field: {e.args[0]}"}), 400
    except (ValueError, TypeError) as e:
        return jsonify({"error": f"Invalid value: {e}"}), 400
    return jsonify(res)


@app.route("/api/predict_file", methods=["POST"])
def api_predict_file():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "No file uploaded"}), 400
    ext = file.filename.lower().rsplit(".", 1)[-1]
    if ext not in ("csv", "xlsx", "xls"):
        return jsonify({"error": "Supported formats: .csv, .xlsx, .xls"}), 400
    path = os.path.join(UPLOAD_DIR, file.filename)
    file.save(path)
    try:
        rows = read_table(path, file.filename)
        rows = normalize_table(rows)
        missing = [c for c in REQUIRED if not any(c in r for r in rows[:1])]
        if missing:
            return jsonify({"error": "Missing columns: " + ", ".join(missing)}), 400
        if not rows:
            return jsonify({"error": "The file contains no data rows"}), 400
        rows, errs = coerce_values(rows)
        proba = predict_batch(rows)
        columns = list(rows[0].keys())
        for i, p in enumerate(proba):
            rows[i]["Churn_Probability"] = round(float(p), 4)
            rows[i]["Churn_Prediction"] = int(p >= THRESHOLD)
            rows[i]["Risk"] = "high" if p >= THRESHOLD else "low"
        columns += ["Churn_Probability", "Churn_Prediction", "Risk"]
        high = sum(1 for r in rows if r["Churn_Prediction"])
        total = len(rows)
        return jsonify({
            "total": total,
            "high_risk": high,
            "high_risk_pct": round(high / total * 100, 2) if total else 0,
            "threshold": THRESHOLD,
            "columns": columns,
            "rows": rows_to_json(rows, columns),
            "errors": errs,
        })
    except Exception as e:
        return jsonify({"error": f"Processing failed: {e}"}), 400
    finally:
        try:
            os.remove(path)
        except OSError:
            pass


@app.route("/api/template")
def template():
    sample = [
        {"Geography": "France", "Gender": "Female", "Age": 42, "Balance": 0,
         "NumOfProducts": 1, "IsActiveMember": 1, "EstimatedSalary": 101348.88},
        {"Geography": "Germany", "Gender": "Male", "Age": 55, "Balance": 120000,
         "NumOfProducts": 2, "IsActiveMember": 0, "EstimatedSalary": 85000.0},
        {"Geography": "Spain", "Gender": "Female", "Age": 33, "Balance": 45000,
         "NumOfProducts": 2, "IsActiveMember": 1, "EstimatedSalary": 92000.0},
    ]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(sample[0].keys()))
    writer.writeheader()
    writer.writerows(sample)
    resp = make_response("\ufeff" + buf.getvalue())
    resp.headers["Content-Disposition"] = "attachment; filename=template.csv"
    resp.headers["Content-Type"] = "text/csv; charset=utf-8"
    return resp


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)